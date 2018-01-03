# SplitExcel.py
# BonizLee
# -*- coding: utf-8 -*-

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
import os.path
import sys

#初始化
def init():
    global PATH
    PATH = os.path.dirname(os.path.realpath(__file__))+os.path.sep
    global EXECLFILE
    EXECLFILE='省厅认定清单20171227.xlsx'
    global FZJG_DIC
    FZJG_DIC={ '粤A':'广州','粤B':'深圳','粤C':'珠海','粤D':'汕头','粤E':'佛山','粤F':'韶关','粤G':'湛江','粤H':'肇庆','粤J':'江门','粤K':'茂名','粤L':'惠州','粤M':'梅州','粤N':'汕尾','粤P':'河源','粤Q':'阳江','粤R':'清远','粤S':'东莞','粤T':'中山','粤U':'潮州','粤V':'揭阳','粤W':'云浮'}

# row_offset表示数据行开始位置，split_column表示拆分列索引
def splitexcel(ws=None,city=None,row_offset=1,split_column=1):
    top_row=next(ws.rows)
    rows_len = ws.max_row
    
    top_row_list = row_to_list(top_row)
    new_wb = Workbook()
    new_ws = new_wb.active
    new_ws.append(top_row_list)
    for row in ws.iter_rows(row_offset=row_offset,max_row=rows_len-1):
        splitcell_value=row[split_column].value       
        splitcell_value=splitcell_value.strip()
        if(splitcell_value==city):
            row_list=row_to_list(row)
            new_ws.append(row_list)
    export_name=xlsx_name(fzjg=city)
    new_wb.save(export_name)
    print(export_name+'完成')


def row_to_list(row=None):
	lis = []
	for cell in row:
		lis.append(cell.value)
	return lis

def xlsx_name(fzjg=None):    
    xlsx_name = FZJG_DIC[fzjg] + '_' +EXECLFILE
    xlsx_name = xlsx_name
    return xlsx_name

if __name__ == "__main__":
    init()
    workbook=load_workbook(PATH+EXECLFILE)
    worksheet=workbook.active
    print('加载配置完成')
    if len(sys.argv)==1:
        for d,v in FZJG_DIC.items():
            splitexcel(ws=worksheet,city=d)

    if len(sys.argv)==2:
        splitexcel(ws=worksheet,city=sys.argv[1])
    
    print('汇总完成完成')
    
    while True:
        q = input("输入Q退出程序：")
        if q == 'Q' or q == 'q':
            break
